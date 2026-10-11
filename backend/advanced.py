"""Office conversion, searchable OCR, and PDF markup engines."""
from contextlib import closing
from functools import lru_cache
import base64
import io
import math
import os
from pathlib import Path
import re
import shutil
import subprocess
import tempfile
import zipfile
from PIL import Image
import pdfplumber
import pypdfium2 as pdfium
import pytesseract
from pypdf import PdfReader, PdfWriter
from reportlab.pdfgen import canvas

ADVANCED = {'word_pdf', 'ppt_pdf', 'excel_pdf', 'pdf_word', 'pdf_ppt', 'pdf_excel', 'ocr', 'edit', 'sign'}
OFFICE_INPUTS = {'word_pdf': {'doc', 'docx'}, 'ppt_pdf': {'ppt', 'pptx'}, 'excel_pdf': {'xls', 'xlsx'}}
RUNTIME = Path(__file__).parent / 'runtime'
LANGUAGES = {'eng', 'ara', 'hin', 'spa', 'fra', 'por', 'ind', 'tur', 'urd'}

@lru_cache(maxsize=4)
def executable(name):
    override = os.getenv('SOFFICE_PATH' if name == 'soffice' else 'TESSERACT_PATH')
    if override and Path(override).is_file():
        return override
    found = shutil.which(name)
    if found:
        return found
    folder = RUNTIME / ('libreoffice' if name == 'soffice' else 'tesseract')
    if folder.exists():
        # Windows soffice.com waits for CLI conversion completion.
        for filename in (['soffice.com', 'soffice.exe'] if name == 'soffice' else ['tesseract.exe']):
            match = next(folder.rglob(filename), None)
            if match:
                return str(match)
    defaults = [Path('C:/Program Files/LibreOffice/program/soffice.com')] if name == 'soffice' else [Path('C:/Program Files/Tesseract-OCR/tesseract.exe')]
    return next((str(p) for p in defaults if p.is_file()), None)


def capability_details():
    return {'office': bool(executable('soffice')), 'ocr': bool(executable('tesseract'))}


def open_pdf(data, password=''):
    reader = PdfReader(io.BytesIO(data), strict=False)
    if reader.is_encrypted and (not password or not reader.decrypt(password)):
        raise ValueError('Unlock this PDF with its correct password first.')
    if not 1 <= len(reader.pages) <= 500:
        raise ValueError('Use a PDF containing between 1 and 500 pages.')
    return reader


def as_bytes(writer):
    out = io.BytesIO()
    writer.write(out)
    return out.getvalue()


def render_page(document, index, scale=2):
    with closing(document[index]) as page:
        w, h = page.get_size()
        if min(w, h) <= 0 or w*h*scale*scale > 30_000_000:
            raise ValueError('Page dimensions are too large to render safely.')
        with closing(page.render(scale=scale)) as bitmap:
            return bitmap.to_pil().copy().convert('RGB')


def office_to_pdf(data, extension):
    engine = executable('soffice')
    if not engine:
        raise ValueError('The processing service needs LibreOffice installed for Office-to-PDF conversion.')
    if extension.endswith('x'):
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as z:
                if len(z.infolist()) > 5000 or sum(i.file_size for i in z.infolist()) > 200*1024*1024:
                    raise ValueError('This Office document expands beyond the processing limit.')
                if '[Content_Types].xml' not in z.namelist():
                    raise ValueError('This is not a valid Office document.')
        except zipfile.BadZipFile:
            raise ValueError('This Office document is damaged or has the wrong extension.') from None
    elif not data.startswith(bytes.fromhex('D0CF11E0A1B11AE1')):
        raise ValueError('This is not a valid legacy Office document.')
    with tempfile.TemporaryDirectory(prefix='ali-docs-office-') as temporary:
        folder = Path(temporary)
        source = folder / f'document.{extension}'
        source.write_bytes(data)
        profile = folder / 'profile'
        (profile / 'user').mkdir(parents=True)
        (profile / 'user/registrymodifications.xcu').write_text('''<?xml version="1.0" encoding="UTF-8"?>
<oor:items xmlns:oor="http://openoffice.org/2001/registry"><item oor:path="/org.openoffice.Office.Common/Security/Scripting"><prop oor:name="MacroSecurityLevel" oor:op="fuse"><value>3</value></prop></item></oor:items>''', encoding='utf-8')
        try:
            subprocess.run([engine, f'-env:UserInstallation={profile.as_uri()}', '--headless', '--nologo', '--nodefault', '--norestore', '--convert-to', 'pdf', '--outdir', str(folder), str(source)],
                           timeout=90, check=True, stdout=subprocess.PIPE, stderr=subprocess.PIPE,
                           creationflags=subprocess.CREATE_NO_WINDOW if os.name == 'nt' else 0)
        except subprocess.TimeoutExpired:
            raise ValueError('Office conversion exceeded 90 seconds. Try a smaller document.') from None
        except subprocess.CalledProcessError:
            raise ValueError('Office conversion failed. Check whether the document opens correctly or needs a password.') from None
        output = folder / 'document.pdf'
        if not output.is_file():
            raise ValueError('LibreOffice could not convert this document. Password-protected Office files are not supported.')
        result = output.read_bytes()
        open_pdf(result)
        return result


def searchable_pdf(data, options):
    reader = open_pdf(data, str(options.get('password', '')))
    engine = executable('tesseract')
    if not engine:
        raise ValueError('The processing service needs Tesseract installed for OCR.')
    pytesseract.pytesseract.tesseract_cmd = engine
    language = str(options.get('language', 'eng'))
    if language not in LANGUAGES:
        raise ValueError('Select a supported OCR language.')
    if language not in pytesseract.get_languages(config=''):
        raise ValueError(f'The {language} OCR language data is not installed on this service.')
    if len(reader.pages) > 40:
        raise ValueError('OCR supports up to 40 pages per task. Split this PDF first.')
    writer = PdfWriter()
    with closing(pdfium.PdfDocument(data, password=str(options.get('password', '')) or None)) as document:
        for i, page in enumerate(reader.pages):
            if (page.extract_text() or '').strip():
                writer.add_page(page)
            else:
                image = render_page(document, i, 3)
                try:
                    page_pdf = pytesseract.image_to_pdf_or_hocr(image, extension='pdf', lang=language, config='--dpi 216', timeout=60)
                except RuntimeError:
                    raise ValueError('OCR timed out on a page. Try a smaller or clearer scan.') from None
                recognized = PdfReader(io.BytesIO(page_pdf))
                writer.add_page(recognized.pages[0])
    return as_bytes(writer)


def readable_source(data, options):
    reader = open_pdf(data, str(options.get('password', '')))
    if any(not (p.extract_text() or '').strip() for p in reader.pages):
        if options.get('ocrFallback', True):
            return searchable_pdf(data, options)
    return data


def pdf_to_office(tool, data, options):
    data = readable_source(data, options)
    password = str(options.get('password', ''))
    reader = open_pdf(data, password)
    if len(reader.pages) > 100:
        raise ValueError('PDF-to-Office conversion supports up to 100 pages per task.')
    with pdfplumber.open(io.BytesIO(data), password=password or None) as pdf:
        if tool == 'pdf_word':
            from docx import Document
            from docx.shared import Inches
            doc = Document()
            first = pdf.pages[0]
            section = doc.sections[0]
            section.page_width, section.page_height = Inches(first.width/72), Inches(first.height/72)
            for i, page in enumerate(pdf.pages):
                if i:
                    doc.add_page_break()
                tables = page.find_tables()
                outside = page.filter(lambda obj: not any(t.bbox[0] <= (obj.get('x0',0)+obj.get('x1',0))/2 <= t.bbox[2] and t.bbox[1] <= (obj.get('top',0)+obj.get('bottom',0))/2 <= t.bbox[3] for t in tables))
                blocks = [(line['top'], 'line', line['text']) for line in outside.extract_text_lines()]
                blocks += [(t.bbox[1], 'table', t.extract()) for t in tables]
                for _, kind, content in sorted(blocks, key=lambda item:item[0]):
                    if kind == 'line':
                        doc.add_paragraph(content)
                    elif content:
                        table = doc.add_table(rows=0, cols=max(map(len,content)))
                        table.style = 'Table Grid'
                        for row in content:
                            cells = table.add_row().cells
                            for c, value in enumerate(row):
                                cells[c].text = str(value or '')
                for image in reader.pages[i].images:
                    try:
                        pixels = image.image.convert('RGB')
                        out = io.BytesIO(); pixels.save(out, 'PNG'); out.seek(0)
                        doc.add_picture(out, width=Inches(min(5.5,first.width/72-1)))
                    except (ValueError, OSError):
                        pass
            out = io.BytesIO(); doc.save(out)
            return out.getvalue(), 'converted.docx', 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
        if tool == 'pdf_excel':
            from openpyxl import Workbook
            workbook = Workbook(); workbook.remove(workbook.active)
            for i, page in enumerate(pdf.pages):
                tables = page.extract_tables()
                if not tables:
                    tables = [[re.split(r'\s{2,}|\t',line) for line in (page.extract_text(layout=True) or '').splitlines() if line.strip()]]
                for j, rows in enumerate(tables):
                    sheet = workbook.create_sheet(f'Page {i+1} Table {j+1}')
                    for r, row in enumerate(rows,1):
                        for c, value in enumerate(row,1):
                            raw = str(value or '').strip()
                            numeric = raw.replace(',', '')
                            if re.fullmatch(r'[+-]?\d+(?:\.\d+)?', numeric) and len(re.sub(r'\D', '', numeric)) <= 15:
                                cell = sheet.cell(r, c, float(numeric) if '.' in numeric else int(numeric))
                            else:
                                cell = sheet.cell(r,c,raw)
                                # Extracted text must never become an executable Excel formula.
                                cell.data_type = 's'
                    sheet.freeze_panes='A2'
                    for col in sheet.columns:
                        sheet.column_dimensions[col[0].column_letter].width=min(60,max(15,max((len(str(c.value or '')) for c in col),default=15)+2))
            if not workbook.sheetnames:
                workbook.create_sheet('Text')
            out=io.BytesIO(); workbook.save(out)
            return out.getvalue(), 'converted.xlsx', 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        if tool == 'pdf_ppt':
            from pptx import Presentation
            from pptx.util import Inches, Pt
            from pptx.dml.color import RGBColor
            presentation=Presentation()
            first=pdf.pages[0]
            presentation.slide_width=Inches(first.width/72)
            presentation.slide_height=Inches(first.height/72)
            editable=options.get('presentationMode','Page images')=='Editable text'
            with closing(pdfium.PdfDocument(data,password=password or None)) as document:
                for i,page in enumerate(pdf.pages):
                    slide=presentation.slides.add_slide(presentation.slide_layouts[6])
                    sx,sy=first.width/page.width,first.height/page.height
                    if not editable:
                        image=render_page(document,i,2)
                        out=io.BytesIO(); image.save(out,'PNG'); out.seek(0)
                        slide.shapes.add_picture(out,0,0,width=presentation.slide_width,height=presentation.slide_height)
                    else:
                        for line in page.extract_text_lines(return_chars=True):
                            if not line['text'].strip(): continue
                            chars=line.get('chars') or [{}]
                            size=float(chars[0].get('size',12))*sy
                            box=slide.shapes.add_textbox(Inches(line['x0']*sx/72),Inches(line['top']*sy/72),Inches(max(20,(line['x1']-line['x0'])*sx+size)/72),Inches(max(size*1.6,(line['bottom']-line['top'])*sy)/72))
                            frame=box.text_frame; frame.margin_left=frame.margin_right=frame.margin_top=frame.margin_bottom=0
                            frame.word_wrap=False
                            run=frame.paragraphs[0].add_run(); run.text=line['text']
                            run.font.size=Pt(size); run.font.name='Arial'; run.font.color.rgb=RGBColor(0,0,0)
                            run.font.bold='bold' in str(chars[0].get('fontname','')).lower()
            out=io.BytesIO(); presentation.save(out)
            return out.getvalue(),'converted.pptx','application/vnd.openxmlformats-officedocument.presentationml.presentation'
    raise ValueError('Unknown PDF conversion tool.')


def apply_markup(data, options, signing=False):
    reader=open_pdf(data,str(options.get('password','')))
    operations=options.get('operations',[])
    if not isinstance(operations,list) or not 1 <= len(operations) <= 200:
        raise ValueError('Add at least one annotation or signature (up to 200).')
    writer=PdfWriter()
    for page in reader.pages:
        writer.add_page(page)
    for operation in operations:
        number=int(operation.get('page',1))
        if not 1 <= number <= len(writer.pages):
            raise ValueError('Annotation page number is outside this document.')
        page=writer.pages[number-1]; page.transfer_rotation_to_content()
        pw,ph=float(page.mediabox.width),float(page.mediabox.height)
        x,y=float(operation.get('x',.1)),float(operation.get('y',.1))
        w,h=float(operation.get('width',.35)),float(operation.get('height',.1))
        if not all(math.isfinite(v) for v in (x,y,w,h)) or not 0 <= x <= 1 or not 0 <= y <= 1 or not 0 < w <= 1 or not 0 < h <= 1 or x+w>1.001 or y+h>1.001:
            raise ValueError('Keep annotations inside the page boundaries.')
        out=io.BytesIO(); overlay=canvas.Canvas(out,pagesize=(pw,ph))
        kind=operation.get('kind','text')
        color=str(operation.get('color','#7428F5'))
        if not re.fullmatch(r'#[0-9a-fA-F]{6}',color): raise ValueError('Invalid annotation color.')
        overlay.setFillColor(color); overlay.setStrokeColor(color)
        px,py=x*pw,(1-y-h)*ph
        if kind in {'text','signature_text'}:
            text=str(operation.get('text','')).strip()
            size=float(operation.get('fontSize',16))
            if not text or len(text)>2000 or not 6 <= size <= 96: raise ValueError('Enter text and select a font size between 6 and 96.')
            overlay.setFont('Helvetica-Oblique' if kind=='signature_text' else 'Helvetica',size)
            for j,line in enumerate(text.splitlines()):
                overlay.drawString(px,(1-y)*ph-size-j*size*1.2,line)
        elif kind in {'signature_image','image'}:
            from reportlab.lib.utils import ImageReader
            encoded=str(operation.get('image',''))
            if len(encoded)>2_000_000: raise ValueError('Use a signature image under 1.5 MB.')
            raw=base64.b64decode(encoded,validate=True)
            image=Image.open(io.BytesIO(raw))
            if image.width*image.height>4_000_000: raise ValueError('Signature image is too large.')
            overlay.drawImage(ImageReader(image),px,py,width=w*pw,height=h*ph,mask='auto',preserveAspectRatio=True,anchor='c')
        elif kind=='highlight':
            overlay.setFillColor('#FFF176'); overlay.setFillAlpha(.35); overlay.rect(px,py,w*pw,h*ph,fill=1,stroke=0)
        elif kind=='whiteout':
            overlay.setFillColor('white'); overlay.rect(px,py,w*pw,h*ph,fill=1,stroke=0)
        elif kind=='rectangle':
            overlay.setLineWidth(2); overlay.rect(px,py,w*pw,h*ph,fill=0,stroke=1)
        else:
            raise ValueError('Unsupported annotation type.')
        overlay.save(); page.merge_page(PdfReader(out).pages[0])
    return as_bytes(writer)


def process_advanced(tool,files,options,filenames=None):
    if len(files)!=1: raise ValueError('This tool accepts one file at a time.')
    if len(files[0])>25*1024*1024: raise ValueError('Use a file under 25 MB.')
    data=files[0]
    if tool in OFFICE_INPUTS:
        filename=(filenames or [f'document.{sorted(OFFICE_INPUTS[tool])[-1]}'])[0]
        extension=Path(filename).suffix.lower().lstrip('.')
        if extension not in OFFICE_INPUTS[tool]: raise ValueError('Choose a supported Office file for this tool.')
        return office_to_pdf(data,extension),'converted.pdf','application/pdf'
    if tool=='ocr': return searchable_pdf(data,options),'searchable.pdf','application/pdf'
    if tool in {'pdf_word','pdf_excel','pdf_ppt'}: return pdf_to_office(tool,data,options)
    if tool in {'edit','sign'}: return apply_markup(data,options,tool=='sign'),f'{tool}.pdf','application/pdf'
    raise ValueError('Unknown processing tool.')
