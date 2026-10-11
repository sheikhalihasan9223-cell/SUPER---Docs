import io
import base64
import tempfile
from pathlib import Path
import pytest
from docx import Document
from openpyxl import Workbook, load_workbook
from pptx import Presentation
from pptx.util import Inches
from PIL import Image, ImageDraw, ImageFont
from reportlab.pdfgen import canvas
from reportlab.lib.utils import ImageReader
from pypdf import PdfReader
from fastapi.testclient import TestClient
from app import app, process
from advanced import executable


def pdf_fixture(table=False, scan=False):
    out=io.BytesIO(); c=canvas.Canvas(out,pagesize=(612,792))
    if scan:
        image=Image.new('RGB',(1200,1600),'white'); d=ImageDraw.Draw(image)
        d.text((110,160),'SEARCHABLE DOCUMENT TEST',font=ImageFont.truetype('C:/Windows/Fonts/arial.ttf',48),fill='black')
        d.text((110,270),'Invoice number 12345',font=ImageFont.truetype('C:/Windows/Fonts/arial.ttf',42),fill='black')
        c.drawImage(ImageReader(image),0,0,width=612,height=792)
    else:
        c.drawString(72,710,'Editable conversion test')
        if table:
            for y in [580,610,640]: c.line(72,y,400,y)
            for x in [72,230,400]: c.line(x,580,x,640)
            c.drawString(82,620,'Item'); c.drawString(240,620,'Amount')
            c.drawString(82,590,'Apples'); c.drawString(240,590,'12')
    c.save(); return out.getvalue()


def office_fixture(kind):
    out=io.BytesIO()
    if kind=='word_pdf':
        doc=Document(); doc.add_heading('Word conversion fixture',0)
        doc.add_paragraph('Office engine works with editable text.')
        doc.save(out)
    elif kind=='excel_pdf':
        workbook=Workbook(); ws=workbook.active
        ws.append(['Item','Amount']); ws.append(['Apples',12]); workbook.save(out)
    else:
        p=Presentation(); s=p.slides.add_slide(p.slide_layouts[6])
        s.shapes.add_textbox(Inches(1),Inches(1),Inches(6),Inches(1)).text='PowerPoint conversion fixture'
        p.save(out)
    return out.getvalue()


@pytest.mark.parametrize('kind,extension,expected',[('word_pdf','docx','Office engine'),('excel_pdf','xlsx','Apples'),('ppt_pdf','pptx','PowerPoint conversion fixture')])
def test_real_office_conversion(kind,extension,expected):
    assert executable('soffice'), 'LibreOffice runtime must be installed'
    data,name,mime=process(kind,[office_fixture(kind)],{},[f'input.{extension}'])
    reader=PdfReader(io.BytesIO(data))
    assert len(reader.pages)>=1
    assert expected in '\n'.join(p.extract_text() or '' for p in reader.pages)
    assert name=='converted.pdf' and mime=='application/pdf'


def test_real_pdf_to_word_text_and_table():
    data,name,_=process('pdf_word',[pdf_fixture(table=True)],{})
    document=Document(io.BytesIO(data))
    assert any('Editable conversion test' in p.text for p in document.paragraphs)
    assert any('Apples' in cell.text for table in document.tables for row in table.rows for cell in row.cells)
    assert name.endswith('.docx')


def test_real_pdf_to_excel_table():
    data,name,_=process('pdf_excel',[pdf_fixture(table=True)],{})
    workbook=load_workbook(io.BytesIO(data))
    values=[c.value for s in workbook for row in s for c in row]
    assert 'Apples' in values and 12 in values
    assert name.endswith('.xlsx')


def test_real_pdf_to_powerpoint_both_modes():
    for mode in ['Page images','Editable text']:
        data,name,_=process('pdf_ppt',[pdf_fixture()],{'presentationMode':mode})
        p=Presentation(io.BytesIO(data)); assert len(p.slides)==1
        if mode=='Page images':
            assert any(s.shape_type==13 for s in p.slides[0].shapes)
        else:
            assert any('Editable conversion test' in getattr(s,'text','') for s in p.slides[0].shapes)
        assert name.endswith('.pptx')


def test_real_searchable_ocr_and_scanned_word():
    assert executable('tesseract'), 'Tesseract runtime must be installed'
    source=pdf_fixture(scan=True)
    data,name,_=process('ocr',[source],{'language':'eng'})
    reader=PdfReader(io.BytesIO(data))
    text=reader.pages[0].extract_text()
    assert '12345' in text
    assert name=='searchable.pdf'
    word,_,_=process('pdf_word',[source],{'language':'eng'})
    doc=Document(io.BytesIO(word))
    assert '12345' in '\n'.join(p.text for p in doc.paragraphs)


def test_annotations_and_signature_modify_real_pdf():
    source=pdf_fixture()
    operations=[{'kind':'text','page':1,'x':.1,'y':.3,'width':.5,'height':.1,'text':'ADDED NOTE','fontSize':16},
                {'kind':'highlight','page':1,'x':.1,'y':.1,'width':.4,'height':.05},
                {'kind':'rectangle','page':1,'x':.1,'y':.4,'width':.4,'height':.1}]
    edited,_,_=process('edit',[source],{'operations':operations})
    assert 'ADDED NOTE' in PdfReader(io.BytesIO(edited)).pages[0].extract_text()
    signature=Image.new('RGBA',(400,100),(0,0,0,0)); d=ImageDraw.Draw(signature)
    d.line([(20,60),(100,20),(160,65),(300,30)],fill='black',width=4)
    png=io.BytesIO();signature.save(png,'PNG')
    signed,_,_=process('sign',[source],{'operations':[{'kind':'signature_image','page':1,'x':.1,'y':.5,'width':.5,'height':.1,'image':base64.b64encode(png.getvalue()).decode()}]})
    assert len(PdfReader(io.BytesIO(signed)).pages[0].images)>=1
    typed,_,_=process('sign',[source],{'operations':[{'kind':'signature_text','page':1,'x':.1,'y':.5,'width':.5,'height':.1,'text':'Ali Hasan','fontSize':22}]})
    assert 'Ali Hasan' in PdfReader(io.BytesIO(typed)).pages[0].extract_text()


def test_invalid_annotations_are_rejected():
    with pytest.raises(ValueError): process('edit',[pdf_fixture()],{'operations':[]})
    with pytest.raises(ValueError): process('sign',[pdf_fixture()],{'operations':[{'kind':'signature_text','page':2,'text':'No'}]})
    with pytest.raises(ValueError): process('edit',[pdf_fixture()],{'operations':[{'kind':'text','x':.9,'width':.5,'text':'No'}]})


def test_office_multipart_contract_and_invalid_input():
    client=TestClient(app)
    response=client.post('/process/word_pdf',files=[('files',('document.docx',office_fixture('word_pdf')))],data={'options':'{}'})
    assert response.status_code==200
    assert 'Office engine' in PdfReader(io.BytesIO(response.content)).pages[0].extract_text()
    assert client.post('/process/word_pdf',files=[('files',('wrong.docx',b'bad'))]).status_code==422
    assert all(name in client.get('/health').json()['tools'] for name in ['ocr','edit','sign','word_pdf','pdf_word','pdf_excel','pdf_ppt'])
